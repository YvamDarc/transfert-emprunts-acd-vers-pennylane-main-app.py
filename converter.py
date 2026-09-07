from __future__ import annotations

import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from statistics import median
from typing import Iterable

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SUMMARY_COLUMNS = [
    "Capital (€)",
    "Taux d’intérêt (%)",
    "Taux d’assurance (%)",
    "Montant total d’assurance (€)",
    "Pourcentage des autres frais (%)",
    "Montant total des autres frais (€)",
    "Nombre d’échéances",
]

SCHEDULE_COLUMNS = [
    "Date",
    "Intérêt (€)",
    "Assurance (€)",
    "Autres frais (€)",
    "Amortissement (€)",
    "Échéance (€)",
    "Solde (€)",
]

SOURCE_COLUMNS = [
    "N°",
    "Date",
    "Annuité",
    "Tx. Int.",
    "Intérêts",
    "Tx. Assur.",
    "Assurance",
    "Tx. Com.",
    "Commission",
    "Capital remb.",
    "Autre",
    "Restant Dû",
]


@dataclass
class ParsedACDLoan:
    data: pd.DataFrame
    source_name: str
    source_type: str
    header_row: int
    warnings: list[str] = field(default_factory=list)
    encoding: str | None = None
    separator: str | None = None


@dataclass
class ConversionResult:
    summary: pd.DataFrame
    schedule: pd.DataFrame
    initial_rows_excluded: int
    zero_installments: int
    frequency_label: str
    source_warnings: list[str] = field(default_factory=list)


def _normalise(text: object) -> str:
    value = "" if text is None else str(text)
    value = unicodedata.normalize("NFKD", value)
    value = "".join(ch for ch in value if not unicodedata.combining(ch))
    value = value.replace("œ", "oe").replace("Œ", "OE")
    value = re.sub(r"[^a-zA-Z0-9]+", " ", value.lower()).strip()
    return re.sub(r"\s+", " ", value)


def _safe_text(value: object) -> str:
    if value is None or pd.isna(value):
        return ""
    return str(value).strip()


def _decode_text(raw: bytes) -> tuple[str, str]:
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    raise ValueError("Le fichier texte ne peut pas être décodé.")


def _number(value: object) -> float | None:
    if value is None or pd.isna(value):
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    text = (
        _safe_text(value)
        .replace("\u00a0", "")
        .replace(" ", "")
        .replace("€", "")
        .replace("%", "")
    )
    if not text:
        return None
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    else:
        text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def _date_value(value: object) -> pd.Timestamp | pd.NaT:
    if value is None or pd.isna(value) or not _safe_text(value):
        return pd.NaT
    if isinstance(value, (pd.Timestamp, datetime, date)):
        return pd.Timestamp(value).normalize()
    if isinstance(value, (int, float)):
        return pd.Timestamp("1899-12-30") + pd.to_timedelta(float(value), unit="D")
    parsed = pd.to_datetime(_safe_text(value), dayfirst=True, errors="coerce")
    return pd.NaT if pd.isna(parsed) else pd.Timestamp(parsed).normalize()


def _is_header(values: Iterable[object]) -> bool:
    normalised = {_normalise(value) for value in values if _safe_text(value)}
    return (
        "date" in normalised
        and any(value.startswith("annuite") for value in normalised)
        and any(value.startswith("interets") or value.startswith("interet") for value in normalised)
        and any("capital remb" in value for value in normalised)
        and any("restant du" in value for value in normalised)
    )


def _find_column(headers: list[object], *aliases: str) -> int | None:
    normalised = [_normalise(value) for value in headers]
    for alias in aliases:
        key = _normalise(alias)
        for index, value in enumerate(normalised):
            if value == key:
                return index
    for alias in aliases:
        key = _normalise(alias)
        for index, value in enumerate(normalised):
            if key and (key in value or value in key):
                return index
    return None


def _cell(row: list[object], index: int | None) -> object:
    if index is None or index >= len(row):
        return ""
    return row[index]


def _rows_to_dataframe(rows: list[list[object]], source_name: str, source_type: str) -> ParsedACDLoan:
    header_index = next((index for index, row in enumerate(rows) if _is_header(row)), None)
    if header_index is None:
        raise ValueError(
            "L'en-tête de l'échéancier ACD n'a pas été reconnu. Les colonnes Date, Annuité, Intérêts, Capital remb. et Restant Dû sont nécessaires."
        )

    headers = rows[header_index]
    indexes = {
        "N°": _find_column(headers, "N°", "No", "Numéro"),
        "Date": _find_column(headers, "Date"),
        "Annuité": _find_column(headers, "Annuité", "Annuite"),
        "Tx. Int.": _find_column(headers, "Tx. Int.", "Taux intérêt", "Taux interet"),
        "Intérêts": _find_column(headers, "Intérêts", "Interets", "Intérêt"),
        "Tx. Assur.": _find_column(headers, "Tx. Assur.", "Taux assurance"),
        "Assurance": _find_column(headers, "Assurance"),
        "Tx. Com.": _find_column(headers, "Tx. Com.", "Taux commission"),
        "Commission": _find_column(headers, "Commission"),
        "Capital remb.": _find_column(headers, "Capital remb.", "Capital remboursé", "Capital rembourse"),
        "Autre": _find_column(headers, "Autre"),
        "Restant Dû": _find_column(headers, "Restant Dû", "Restant Du", "Capital restant dû", "Solde"),
    }
    missing = [name for name in ("Date", "Annuité", "Intérêts", "Capital remb.", "Restant Dû") if indexes[name] is None]
    if missing:
        raise ValueError("Colonnes ACD obligatoires absentes : " + ", ".join(missing))

    records: list[dict[str, object]] = []
    for row in rows[header_index + 1 :]:
        if not any(_safe_text(value) for value in row):
            continue
        row_date = _date_value(_cell(row, indexes["Date"]))
        if pd.isna(row_date):
            continue
        record: dict[str, object] = {
            "N°": _number(_cell(row, indexes["N°"])),
            "Date": row_date,
        }
        for column in SOURCE_COLUMNS[2:]:
            record[column] = _number(_cell(row, indexes[column]))
        records.append(record)

    if not records:
        raise ValueError("Aucune échéance datée n'a été détectée dans le fichier ACD.")

    data = pd.DataFrame(records, columns=SOURCE_COLUMNS)
    return ParsedACDLoan(
        data=data,
        source_name=source_name,
        source_type=source_type,
        header_row=header_index + 1,
    )


def _parse_text(source_name: str, raw: bytes, source_type: str = "CSV / texte") -> ParsedACDLoan:
    text, encoding = _decode_text(raw)
    lines = [line for line in text.splitlines() if line.strip()]
    header_line = next((line for line in lines if "capital" in _normalise(line) and "restant" in _normalise(line)), "")
    candidates = ("\t", ";", "|")
    separator = max(candidates, key=lambda item: header_line.count(item))
    if header_line.count(separator) < 2:
        try:
            separator = csv.Sniffer().sniff(text[:10000], delimiters="\t;|").delimiter
        except csv.Error as exc:
            raise ValueError("Le séparateur n'a pas été reconnu. Copiez le tableau avec ses tabulations depuis ACD ou Excel.") from exc
    rows = list(csv.reader(io.StringIO(text), delimiter=separator))
    parsed = _rows_to_dataframe(rows, source_name, source_type)
    parsed.encoding = encoding
    parsed.separator = separator
    return parsed


def _parse_excel(source_name: str, raw: bytes) -> ParsedACDLoan:
    workbook = pd.ExcelFile(io.BytesIO(raw))
    for sheet_name in workbook.sheet_names:
        frame = pd.read_excel(workbook, sheet_name=sheet_name, header=None, dtype=object)
        rows = frame.values.tolist()
        if any(_is_header(row) for row in rows[:200]):
            return _rows_to_dataframe(rows, source_name, f"Excel - {sheet_name}")
    raise ValueError("Aucun onglet ne contient l'échéancier ACD attendu.")


def parse_acd_loan_file(source_name: str, raw: bytes) -> ParsedACDLoan:
    suffix = Path(source_name).suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return _parse_excel(source_name, raw)
    if suffix in {".csv", ".txt", ".tsv", ".asc"}:
        return _parse_text(source_name, raw)
    raise ValueError("Format non pris en charge. Déposez un CSV, TXT, TSV ou XLSX ACD.")


def parse_acd_loan_paste(text: str) -> ParsedACDLoan:
    return _parse_text("copier_coller_acd.txt", text.encode("utf-8"), source_type="Copier-coller")


def _frequency(dates: pd.Series) -> tuple[int, str]:
    clean_dates = sorted(pd.Timestamp(value) for value in dates.dropna().unique())
    if len(clean_dates) < 2:
        return 12, "mensuelle (estimée)"
    gaps = [(later - earlier).days for earlier, later in zip(clean_dates, clean_dates[1:]) if later > earlier]
    if not gaps:
        return 12, "mensuelle (estimée)"
    typical = median(gaps)
    if typical <= 45:
        return 12, "mensuelle"
    if typical <= 120:
        return 4, "trimestrielle"
    if typical <= 220:
        return 2, "semestrielle"
    return 1, "annuelle"


def _sum_columns(frame: pd.DataFrame, columns: list[str]) -> pd.Series:
    return frame[columns].fillna(0).astype(float).sum(axis=1)


def convert_acd_loan(
    parsed: ParsedACDLoan,
    payment_method: str = "annuity_plus_fees",
    keep_zero_installments: bool = True,
) -> ConversionResult:
    source = parsed.data.copy().reset_index(drop=True)
    money_columns = ["Annuité", "Intérêts", "Assurance", "Commission", "Capital remb.", "Autre"]
    for column in money_columns + ["Restant Dû", "Tx. Int.", "Tx. Assur.", "Tx. Com."]:
        source[column] = pd.to_numeric(source[column], errors="coerce")

    zero_flow = _sum_columns(source, money_columns).abs() <= 0.005
    numbered_zero = pd.to_numeric(source["N°"], errors="coerce").fillna(-1).eq(0)
    initial_mask = numbered_zero & zero_flow & source["Restant Dû"].notna()
    initial_rows_excluded = int(initial_mask.sum())

    if initial_mask.any():
        capital = float(source.loc[initial_mask, "Restant Dû"].iloc[0])
    else:
        first = source.iloc[0]
        remaining = _number(first.get("Restant Dû")) or 0.0
        principal = _number(first.get("Capital remb.")) or 0.0
        capital = float(remaining + principal)

    schedule_source = source.loc[~initial_mask].copy().reset_index(drop=True)
    zero_schedule = _sum_columns(schedule_source, money_columns).abs() <= 0.005
    zero_installments = int(zero_schedule.sum())
    if not keep_zero_installments:
        schedule_source = schedule_source.loc[~zero_schedule].copy().reset_index(drop=True)

    other_fees = schedule_source["Commission"].fillna(0) + schedule_source["Autre"].fillna(0)
    detailed_payment = (
        schedule_source["Intérêts"].fillna(0)
        + schedule_source["Assurance"].fillna(0)
        + other_fees
        + schedule_source["Capital remb."].fillna(0)
    )
    if payment_method == "components":
        payments = detailed_payment
    else:
        payments = schedule_source["Annuité"].fillna(0) + schedule_source["Assurance"].fillna(0) + other_fees

    schedule = pd.DataFrame(
        {
            "Date": pd.to_datetime(schedule_source["Date"], errors="coerce"),
            "Intérêt (€)": schedule_source["Intérêts"].fillna(0).round(2),
            "Assurance (€)": schedule_source["Assurance"].fillna(0).round(2),
            "Autres frais (€)": other_fees.round(2),
            "Amortissement (€)": schedule_source["Capital remb."].fillna(0).round(2),
            "Échéance (€)": payments.round(2),
            "Solde (€)": schedule_source["Restant Dû"].round(2),
        },
        columns=SCHEDULE_COLUMNS,
    )

    periods_per_year, frequency_label = _frequency(schedule["Date"])
    positive_period_rates = schedule_source.loc[schedule_source["Tx. Int."] > 0, "Tx. Int."].dropna()
    annual_rate = float(positive_period_rates.median() * periods_per_year) if len(positive_period_rates) else 0.0
    total_insurance = float(schedule["Assurance (€)"].sum())
    total_other_fees = float(schedule["Autres frais (€)"].sum())
    insurance_pct = total_insurance / capital * 100 if capital else 0.0
    other_fees_pct = total_other_fees / capital * 100 if capital else 0.0

    summary = pd.DataFrame(
        [
            {
                "Capital (€)": round(capital, 2),
                "Taux d’intérêt (%)": round(annual_rate, 6),
                "Taux d’assurance (%)": round(insurance_pct, 6),
                "Montant total d’assurance (€)": 0.0,
                "Pourcentage des autres frais (%)": round(other_fees_pct, 6),
                "Montant total des autres frais (€)": 0.0,
                "Nombre d’échéances": len(schedule),
            }
        ],
        columns=SUMMARY_COLUMNS,
    )

    source_warnings: list[str] = []
    annuity_difference = (
        schedule_source["Annuité"].fillna(0)
        - schedule_source["Intérêts"].fillna(0)
        - schedule_source["Capital remb."].fillna(0)
    ).abs()
    inconsistent = int((annuity_difference > 0.02).sum())
    if inconsistent:
        source_warnings.append(
            f"{inconsistent} ligne(s) présentent un écart entre l'annuité ACD et intérêts + capital remboursé. Vérifiez le mode de calcul de l'échéance."
        )

    return ConversionResult(
        summary=summary,
        schedule=schedule,
        initial_rows_excluded=initial_rows_excluded,
        zero_installments=zero_installments,
        frequency_label=frequency_label,
        source_warnings=source_warnings,
    )


def validate_pennylane_loan(summary: pd.DataFrame, schedule: pd.DataFrame) -> pd.DataFrame:
    issues: list[dict[str, object]] = []

    def add(level: str, row: str | int, field: str, message: str) -> None:
        issues.append({"Niveau": level, "Ligne": row, "Champ": field, "Message": message})

    if summary.empty:
        add("Erreur", "Paramètres", "Synthèse", "La ligne de paramètres est absente.")
        return pd.DataFrame(issues, columns=["Niveau", "Ligne", "Champ", "Message"])

    params = summary.iloc[0]
    capital = _number(params.get("Capital (€)"))
    count = _number(params.get("Nombre d’échéances"))
    if capital is None or capital <= 0:
        add("Erreur", "Paramètres", "Capital (€)", "Le capital initial doit être supérieur à zéro.")
    if count is None or int(round(count)) != len(schedule):
        add("Erreur", "Paramètres", "Nombre d’échéances", f"La valeur doit être égale au nombre de lignes : {len(schedule)}.")
    for field in [
        "Taux d’intérêt (%)",
        "Taux d’assurance (%)",
        "Montant total d’assurance (€)",
        "Pourcentage des autres frais (%)",
        "Montant total des autres frais (€)",
    ]:
        value = _number(params.get(field))
        if value is None or value < 0:
            add("Erreur", "Paramètres", field, "La valeur doit être numérique et positive ou nulle.")

    if schedule.empty:
        add("Erreur", "Échéancier", "Lignes", "Aucune échéance à importer.")
        return pd.DataFrame(issues, columns=["Niveau", "Ligne", "Champ", "Message"])

    previous_date: pd.Timestamp | None = None
    previous_balance = capital
    for index, row in schedule.reset_index(drop=True).iterrows():
        excel_row = index + 6
        row_date = _date_value(row.get("Date"))
        if pd.isna(row_date):
            add("Erreur", excel_row, "Date", "Date absente ou illisible.")
        elif previous_date is not None and row_date <= previous_date:
            add("Erreur", excel_row, "Date", "Les dates doivent être strictement croissantes.")
        if not pd.isna(row_date):
            previous_date = row_date

        amounts: dict[str, float | None] = {}
        for field in SCHEDULE_COLUMNS[1:]:
            amounts[field] = _number(row.get(field))
            if amounts[field] is None:
                add("Erreur", excel_row, field, "Montant absent ou illisible.")
            elif amounts[field] < -0.005:
                add("Avertissement", excel_row, field, "Montant négatif à confirmer.")

        components = [amounts.get(field) for field in ["Intérêt (€)", "Assurance (€)", "Autres frais (€)", "Amortissement (€)"]]
        payment = amounts.get("Échéance (€)")
        if all(value is not None for value in components) and payment is not None:
            expected_payment = sum(float(value) for value in components if value is not None)
            if abs(payment - expected_payment) > 0.02:
                add("Erreur", excel_row, "Échéance (€)", f"L'échéance diffère de la somme des composantes de {payment - expected_payment:.2f} €.")

        principal = amounts.get("Amortissement (€)")
        balance = amounts.get("Solde (€)")
        if previous_balance is not None and principal is not None and balance is not None:
            expected_balance = previous_balance - principal
            if abs(balance - expected_balance) > 0.05:
                add("Erreur", excel_row, "Solde (€)", f"Solde attendu : {expected_balance:.2f} €.")
            previous_balance = balance

    if previous_balance is not None and previous_balance > 0.05:
        add("Avertissement", "Fin", "Solde (€)", f"Le dernier solde n'est pas nul : {previous_balance:.2f} €.")

    return pd.DataFrame(issues, columns=["Niveau", "Ligne", "Champ", "Message"])


def _excel_value(value: object) -> object:
    if value is None or pd.isna(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if isinstance(value, date) and not isinstance(value, datetime):
        return datetime.combine(value, datetime.min.time())
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def to_xlsx_bytes(summary: pd.DataFrame, schedule: pd.DataFrame) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for column_index, header in enumerate(SUMMARY_COLUMNS, start=1):
        cell = sheet.cell(row=1, column=column_index, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        value = _excel_value(summary.iloc[0][header])
        sheet.cell(row=2, column=column_index, value=value)

    for column_index, header in enumerate(SCHEDULE_COLUMNS, start=1):
        cell = sheet.cell(row=5, column=column_index, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_index, (_, row) in enumerate(schedule[SCHEDULE_COLUMNS].iterrows(), start=6):
        for column_index, header in enumerate(SCHEDULE_COLUMNS, start=1):
            sheet.cell(row=row_index, column=column_index, value=_excel_value(row[header]))

    for cell in sheet[2]:
        if cell.column in {1, 4, 6}:
            cell.number_format = "#,##0.00"
        elif cell.column in {2, 3, 5}:
            cell.number_format = "0.000000"
        elif cell.column == 7:
            cell.number_format = "0"
    for row in sheet.iter_rows(min_row=6, max_row=sheet.max_row, min_col=1, max_col=7):
        row[0].number_format = "dd/mm/yyyy"
        for cell in row[1:]:
            cell.number_format = "#,##0.00"

    widths = [18, 20, 21, 31, 35, 34, 20]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 42
    sheet.row_dimensions[5].height = 28
    sheet.freeze_panes = "A6"

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _formatted(value: object, column: str) -> str:
    if value is None or pd.isna(value):
        return ""
    if column == "Date":
        parsed = _date_value(value)
        return "" if pd.isna(parsed) else parsed.strftime("%d/%m/%Y")
    number = _number(value)
    if number is None:
        return _safe_text(value).replace("\t", " ").replace("\n", " ")
    if column == "Nombre d’échéances":
        return str(int(round(number)))
    text = f"{number:.6f}".rstrip("0").rstrip(".")
    return text.replace(".", ",")


def to_full_rows(summary: pd.DataFrame, schedule: pd.DataFrame) -> list[list[str]]:
    rows: list[list[str]] = []
    rows.append(SUMMARY_COLUMNS)
    rows.append([_formatted(summary.iloc[0][column], column) for column in SUMMARY_COLUMNS])
    rows.append([""] * 7)
    rows.append([""] * 7)
    rows.append(SCHEDULE_COLUMNS)
    for _, row in schedule[SCHEDULE_COLUMNS].iterrows():
        rows.append([_formatted(row[column], column) for column in SCHEDULE_COLUMNS])
    return rows


def to_csv_bytes(summary: pd.DataFrame, schedule: pd.DataFrame, separator: str = ";") -> bytes:
    output = io.StringIO(newline="")
    writer = csv.writer(output, delimiter=separator, quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")
    writer.writerows(to_full_rows(summary, schedule))
    return ("\ufeff" + output.getvalue()).encode("utf-8")


def to_clipboard_tsv(summary: pd.DataFrame, schedule: pd.DataFrame) -> str:
    return "\n".join("\t".join(row) for row in to_full_rows(summary, schedule))
